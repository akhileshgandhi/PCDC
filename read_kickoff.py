import docx
p = r"C:\Users\user\Downloads\PCDC Project Kick-off Meeting .docx"
d = docx.Document(p)
out = []
for para in d.paragraphs:
    t = para.text.strip()
    if t:
        out.append(t)
for tbl in d.tables:
    out.append("--- TABLE ---")
    for row in tbl.rows:
        out.append(" | ".join(c.text.strip().replace("\n"," ") for c in row.cells))
text = "\n".join(out)
with open(r"C:\Users\user\Downloads\Prestige\kickoff_extracted.txt","w",encoding="utf-8") as f:
    f.write(text)
print(f"chars: {len(text)}  lines: {len(out)}")
